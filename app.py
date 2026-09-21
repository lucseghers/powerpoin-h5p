import streamlit as st
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw, ImageFont
from pptx.dml.color import RGBColor
import tempfile
import zipfile
import shutil
import uuid
import json
import html
import io
import os


st.set_page_config(
    page_title="PowerPoint naar H5P",
    page_icon="📊",
    layout="centered",
)


APP_DIR = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = APP_DIR / "templates" / "default_template.h5p"


# ---------------------------------------------------------
# Hulpfuncties
# ---------------------------------------------------------

def pct(value, total):
    return max(0, min(100, value / total * 100))


def run_html(run):
    text = html.escape(run.text).replace("\n", "<br>")
    if not text:
        return ""

    styles = []
    font = run.font

    if font.size:
        styles.append(f"font-size:{font.size.pt:.1f}pt")

    try:
        if font.color and font.color.type is not None and font.color.rgb:
            styles.append(f"color:#{font.color.rgb}")
    except Exception:
        pass

    if font.bold:
        text = f"<strong>{text}</strong>"
    if font.italic:
        text = f"<em>{text}</em>"
    if font.underline:
        text = f'<span style="text-decoration:underline">{text}</span>'

    if styles:
        text = f'<span style="{";".join(styles)}">{text}</span>'

    return text


def paragraph_html(p):
    body = "".join(run_html(r) for r in p.runs)

    if not body:
        body = html.escape(p.text)

    styles = []

    if p.alignment == PP_ALIGN.CENTER:
        styles.append("text-align:center")
    elif p.alignment == PP_ALIGN.RIGHT:
        styles.append("text-align:right")
    elif p.alignment == PP_ALIGN.JUSTIFY:
        styles.append("text-align:justify")

    style = f' style="{";".join(styles)}"' if styles else ""
    return f"<p{style}>{body}</p>"


def text_frame_html(tf):
    return "".join(
        paragraph_html(p)
        for p in tf.paragraphs
        if p.text.strip()
    )


def common_element_fields(x, y, w, h):
    return {
        "x": x,
        "y": y,
        "width": w,
        "height": h,
        "alwaysDisplayComments": False,
        "backgroundOpacity": 0,
        "displayAsButton": False,
        "buttonSize": "big",
        "goToSlideType": "specified",
        "invisible": False,
        "solution": "",
    }


def make_text_element(shape, slide_w, slide_h):
    x = pct(shape.left, slide_w)
    y = pct(shape.top, slide_h)
    w = pct(shape.width, slide_w)
    h = pct(shape.height, slide_h)

    text = text_frame_html(shape.text_frame)

    element = common_element_fields(x, y, w, h)

    element["action"] = {
        "library": "H5P.AdvancedText 1.1",
        "params": {
            "text": text
        },
        "subContentId": str(uuid.uuid4()),
        "metadata": {
            "contentType": "Text",
            "license": "U",
            "title": shape.name or "PowerPoint tekst",
            "authors": [],
            "changes": [],
        },
    }

    return element



def _rgb(color, fallback):
    """Lees een expliciete PowerPoint RGB-kleur; themakleuren krijgen een fallback."""
    try:
        if color is not None and color.rgb is not None:
            return tuple(bytes(color.rgb))
    except (AttributeError, TypeError, ValueError):
        pass
    return fallback


def _font(size, bold=False):
    """Zoek een schaalbaar standaardlettertype dat op de server aanwezig is."""
    from functools import lru_cache
    return _system_font(max(9, int(size)), bold)


@__import__("functools").lru_cache(maxsize=64)
def _system_font(size, bold=False):
    # Streamlit Community Cloud kan andere fonts hebben dan een lokale computer.
    # Probeer gangbare systeemfonts, zonder een specifiek font te vereisen.
    names = (
        ["LiberationSans-Bold.ttf", "DejaVuSans-Bold.ttf", "Lato-Bold.ttf",
         "NotoSans-Bold.ttf", "Arial-Bold.ttf"] if bold else
        ["LiberationSans-Regular.ttf", "DejaVuSans.ttf", "Lato-Regular.ttf",
         "NotoSans-Regular.ttf", "Arial.ttf"]
    )
    roots = [Path("/usr/share/fonts"), Path("/usr/local/share/fonts"),
             Path.home() / ".fonts", Path.home() / ".local/share/fonts"]
    for name in names:
        try:
            return ImageFont.truetype(name, size)
        except (OSError, ValueError):
            pass
        for root in roots:
            if root.is_dir():
                for candidate in root.rglob(name):
                    try:
                        return ImageFont.truetype(str(candidate), size)
                    except (OSError, ValueError):
                        pass
    # Laatste uitweg: ieder beschikbaar schaalbaar font (nooit load_default,
    # want dat kan een klein bitmapfont zijn dat de gevraagde grootte negeert).
    for root in roots:
        if root.is_dir():
            for pattern in ("*.ttf", "*.otf"):
                for candidate in root.rglob(pattern):
                    try:
                        return ImageFont.truetype(str(candidate), size)
                    except (OSError, ValueError):
                        pass
    raise RuntimeError(
        "Er is geen schaalbaar systeemlettertype (.ttf/.otf) gevonden op de "
        "Streamlit-server. Voeg een font toe aan de serveromgeving."
    )


def _wrap_text(draw, text, font, max_width):
    lines = []
    for paragraph in text.splitlines() or [""]:
        words = paragraph.split()
        if not words:
            lines.append("")
            continue
        line = ""
        for word in words:
            candidate = f"{line} {word}" if line else word
            if line and draw.textlength(candidate, font=font) > max_width:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    return lines


def render_table_png(shape, images_dir, slide_no, image_no):
    """Tabel als PNG met leesbare, vaste tekstgrootte en regelafbreking."""
    table = shape.table
    scale = 150 / 914400
    widths = [max(1, round(col.width * scale)) for col in table.columns]
    heights = [max(1, round(row.height * scale)) for row in table.rows]
    image = Image.new("RGB", (sum(widths), sum(heights)), "white")
    draw = ImageDraw.Draw(image)
    y = 0
    for row_idx, row in enumerate(table.rows):
        x = 0
        for col_idx, cell in enumerate(row.cells):
            w, h = widths[col_idx], heights[row_idx]
            try:
                fill = _rgb(cell.fill.fore_color, (255, 255, 255))
            except (AttributeError, TypeError, ValueError):
                fill = (255, 255, 255)
            draw.rectangle((x, y, x + w - 1, y + h - 1),
                           fill=fill, outline=(230, 235, 235), width=1)
            paragraph = next((p for p in cell.text_frame.paragraphs if p.text.strip()), None)
            run = next((r for r in paragraph.runs if r.text.strip()), None) if paragraph else None
            bold = row_idx == 0 or (bool(run.font.bold) if run else False)
            foreground = _rgb(run.font.color if run else None,
                              (255, 255, 255) if row_idx == 0 else (45, 52, 55))
            # Op 150 dpi wordt 30 px ongeveer 22 px op een H5P-dia van 1100 px breed.
            # Geen automatische verkleining tot 10 px meer.
            font_px = 32 if row_idx == 0 else 30
            font = _font(font_px, bold)
            pad_x, pad_y = 12, 5
            lines = _wrap_text(draw, cell.text, font, max(1, w - 2 * pad_x))
            line_height = font_px + 4
            required = len(lines) * line_height
            if required > h - 2 * pad_y:
                # Bewaar de tekst: in een uitzonderlijk volle cel alleen tot 24 px verkleinen.
                font_px = 24
                font = _font(font_px, bold)
                lines = _wrap_text(draw, cell.text, font, max(1, w - 2 * pad_x))
                line_height = font_px + 3
                required = len(lines) * line_height
            text_y = y + max(pad_y, (h - required) // 2)
            for line in lines:
                draw.text((x + pad_x, text_y), line, font=font, fill=foreground)
                text_y += line_height
            x += w
        y += heights[row_idx]
    name = f"ppt_s{slide_no:03d}_table{image_no:03d}.png"
    image.save(images_dir / name, format="PNG")
    return name, image.width, image.height


def make_table_element(shape, slide_w, slide_h, images_dir, slide_no, image_no):
    """Plaats de gerasterde tabel als H5P.Image op de oorspronkelijke positie."""
    name, px_w, px_h = render_table_png(shape, images_dir, slide_no, image_no)
    element = common_element_fields(
        pct(shape.left, slide_w), pct(shape.top, slide_h),
        pct(shape.width, slide_w), pct(shape.height, slide_h)
    )
    element["action"] = {
        "library": "H5P.Image 1.1",
        "params": {
            "decorative": True,
            "contentName": "Image",
            "expandImage": "Expand Image",
            "minimizeImage": "Minimize Image",
            "file": {
                "path": f"images/{name}", "mime": "image/png",
                "copyright": {"license": "U"},
                "width": px_w, "height": px_h,
            },
        },
        "subContentId": str(uuid.uuid4()),
        "metadata": {
            "contentType": "Image", "license": "U",
            "title": shape.name or "PowerPoint tabel als afbeelding",
            "authors": [], "changes": [],
        },
    }
    return element


def make_image_element(
    shape,
    slide_w,
    slide_h,
    images_dir,
    slide_no,
    image_no,
):
    image = shape.image
    ext = image.ext.lower()

    if ext == "jpeg":
        ext = "jpg"

    name = f"ppt_s{slide_no:03d}_img{image_no:03d}.{ext}"
    target = images_dir / name
    target.write_bytes(image.blob)

    try:
        with Image.open(target) as im:
            px_w, px_h = im.size
    except Exception:
        px_w, px_h = 0, 0

    mime_map = {
        "png": "image/png",
        "jpg": "image/jpeg",
        "gif": "image/gif",
        "bmp": "image/bmp",
        "tif": "image/tiff",
        "tiff": "image/tiff",
    }

    mime = mime_map.get(ext, f"image/{ext}")

    x = pct(shape.left, slide_w)
    y = pct(shape.top, slide_h)
    w = pct(shape.width, slide_w)
    h = pct(shape.height, slide_h)

    params = {
        "decorative": True,
        "contentName": "Image",
        "expandImage": "Expand Image",
        "minimizeImage": "Minimize Image",
        "file": {
            "path": f"images/{name}",
            "mime": mime,
            "copyright": {
                "license": "U"
            },
            "width": px_w,
            "height": px_h,
        },
    }

    element = common_element_fields(x, y, w, h)

    element["action"] = {
        "library": "H5P.Image 1.1",
        "params": params,
        "subContentId": str(uuid.uuid4()),
        "metadata": {
            "contentType": "Image",
            "license": "U",
            "title": shape.name or "PowerPoint afbeelding",
            "authors": [],
            "changes": [],
        },
    }

    return element


def convert_pptx_to_h5p(pptx_bytes, template_bytes, pptx_name):
    workdir = Path(tempfile.mkdtemp(prefix="ppt_streamlit_h5p_"))

    try:
        pptx_path = workdir / "input.pptx"
        template_path = workdir / "template.h5p"
        h5p_dir = workdir / "h5p"

        pptx_path.write_bytes(pptx_bytes)
        template_path.write_bytes(template_bytes)

        h5p_dir.mkdir(parents=True, exist_ok=True)

        # H5P-template uitpakken
        with zipfile.ZipFile(template_path, "r") as z:
            z.extractall(h5p_dir)

        content_path = h5p_dir / "content" / "content.json"
        meta_path = h5p_dir / "h5p.json"
        images_dir = h5p_dir / "content" / "images"
        images_dir.mkdir(parents=True, exist_ok=True)

        if not content_path.exists():
            raise ValueError(
                "Het gekozen H5P-bestand bevat geen content/content.json."
            )

        if not meta_path.exists():
            raise ValueError(
                "Het gekozen H5P-bestand bevat geen h5p.json."
            )

        with open(content_path, "r", encoding="utf-8") as f:
            content = json.load(f)

        with open(meta_path, "r", encoding="utf-8") as f:
            meta = json.load(f)

        if "presentation" not in content:
            raise ValueError(
                "Het gekozen H5P-template lijkt geen Course Presentation te zijn."
            )

        prs = Presentation(pptx_path)
        slide_w = prs.slide_width
        slide_h = prs.slide_height

        new_slides = []
        stats = {
            "slides": len(prs.slides),
            "tekst": 0,
            "afbeelding": 0,
            "tabel": 0,
            "overgeslagen": 0,
        }

        warnings = []

        for slide_no, slide in enumerate(prs.slides, start=1):
            elements = []
            image_no = 0
            table_no = 0

            # python-pptx geeft shapes in z-volgorde terug.
            for shape in slide.shapes:
                try:
                    if getattr(shape, "has_table", False):
                        table_no += 1
                        elements.append(
                            make_table_element(
                                shape,
                                slide_w,
                                slide_h,
                                images_dir,
                                slide_no,
                                table_no,
                            )
                        )
                        stats["tabel"] += 1

                    elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        image_no += 1
                        elements.append(
                            make_image_element(
                                shape,
                                slide_w,
                                slide_h,
                                images_dir,
                                slide_no,
                                image_no,
                            )
                        )
                        stats["afbeelding"] += 1

                    elif (
                        getattr(shape, "has_text_frame", False)
                        and shape.text.strip()
                    ):
                        elements.append(
                            make_text_element(
                                shape,
                                slide_w,
                                slide_h,
                            )
                        )
                        stats["tekst"] += 1

                    else:
                        stats["overgeslagen"] += 1

                except Exception as exc:
                    stats["overgeslagen"] += 1
                    warnings.append(
                        f"Dia {slide_no} – {shape.name}: {exc}"
                    )

            new_slides.append(
                {
                    "elements": elements,
                    "keywords": [],
                    "slideBackgroundSelector": {},
                }
            )

        content["presentation"]["slides"] = new_slides

        title = Path(pptx_name).stem
        meta["title"] = title
        meta["extraTitle"] = title

        with open(content_path, "w", encoding="utf-8") as f:
            json.dump(
                content,
                f,
                ensure_ascii=False,
                separators=(",", ":"),
            )

        with open(meta_path, "w", encoding="utf-8") as f:
            json.dump(
                meta,
                f,
                ensure_ascii=False,
                separators=(",", ":"),
            )

        output_buffer = io.BytesIO()

        with zipfile.ZipFile(
            output_buffer,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as z:
            for file in h5p_dir.rglob("*"):
                if file.is_file():
                    z.write(
                        file,
                        file.relative_to(h5p_dir).as_posix(),
                    )

        output_buffer.seek(0)

        output_name = f"{title}_bewerkbaar_H5P.h5p"

        return output_buffer.getvalue(), output_name, stats, warnings

    finally:
        shutil.rmtree(workdir, ignore_errors=True)


# ---------------------------------------------------------
# Interface
# ---------------------------------------------------------

st.title("PowerPoint → H5P Course Presentation")
st.caption("Versie 1.5 – automatisch systeemlettertype")

st.write(
    "Zet een PowerPoint om naar een bewerkbare H5P Course Presentation. "
    "Tekstvakken en afbeeldingen blijven afzonderlijke H5P-elementen; tabellen worden als PNG overgenomen."
)

st.info(
    "Momenteel worden tekst en afbeeldingen ondersteund; tabellen worden afbeeldingen. "
    "SmartArt, grafieken, vormen en animaties worden voorlopig overgeslagen."
)

st.subheader("1. PowerPoint")

ppt_file = st.file_uploader(
    "Kies een PowerPoint-bestand",
    type=["pptx"],
    help="Upload de PowerPoint die je naar H5P wilt omzetten.",
)

st.subheader("2. H5P-template")

use_custom_template = st.checkbox(
    "Een ander H5P-template gebruiken",
    value=False,
)

custom_template = None

if use_custom_template:
    custom_template = st.file_uploader(
        "Kies een H5P Course Presentation als template",
        type=["h5p"],
        help=(
            "De achtergrond en instellingen van dit template "
            "worden gebruikt voor de nieuwe presentatie."
        ),
    )
else:
    if DEFAULT_TEMPLATE.exists():
        st.success("Standaardtemplate geladen: default_template.h5p")
    else:
        st.error(
            "Het standaardtemplate ontbreekt in de map templates."
        )


st.subheader("3. Omzetten")

can_convert = (
    ppt_file is not None
    and (
        (not use_custom_template and DEFAULT_TEMPLATE.exists())
        or
        (use_custom_template and custom_template is not None)
    )
)

if st.button(
    "PowerPoint omzetten naar H5P",
    type="primary",
    disabled=not can_convert,
    use_container_width=True,
):

    with st.spinner("PowerPoint wordt omgezet..."):
        try:
            pptx_bytes = ppt_file.getvalue()

            if use_custom_template:
                template_bytes = custom_template.getvalue()
                gebruikte_template = custom_template.name
            else:
                template_bytes = DEFAULT_TEMPLATE.read_bytes()
                gebruikte_template = "default_template.h5p"

            (
                h5p_bytes,
                output_name,
                stats,
                warnings,
            ) = convert_pptx_to_h5p(
                pptx_bytes,
                template_bytes,
                ppt_file.name,
            )

            st.session_state["h5p_bytes"] = h5p_bytes
            st.session_state["output_name"] = output_name
            st.session_state["stats"] = stats
            st.session_state["warnings"] = warnings
            st.session_state["template_name"] = gebruikte_template

        except Exception as exc:
            st.exception(exc)


if "h5p_bytes" in st.session_state:

    st.success("De H5P Course Presentation is aangemaakt.")

    stats = st.session_state["stats"]

    col1, col2, col3, col4, col5 = st.columns(5)

    col1.metric("Dia's", stats["slides"])
    col2.metric("Tekst", stats["tekst"])
    col3.metric("Afbeeldingen", stats["afbeelding"])
    col4.metric("Tabellen", stats["tabel"])
    col5.metric("Overgeslagen", stats["overgeslagen"])

    st.caption(
        f"Gebruikt template: {st.session_state['template_name']}"
    )

    warnings = st.session_state.get("warnings", [])

    if warnings:
        with st.expander(
            f"Waarschuwingen bekijken ({len(warnings)})"
        ):
            for warning in warnings:
                st.write("•", warning)

    st.download_button(
        label="Download H5P-bestand",
        data=st.session_state["h5p_bytes"],
        file_name=st.session_state["output_name"],
        mime="application/octet-stream",
        use_container_width=True,
        type="primary",
    )


st.divider()

st.caption(
    "Standaardtemplate wijzigen: vervang "
    "`templates/default_template.h5p` door een ander getest "
    "H5P Course Presentation-template."
)
